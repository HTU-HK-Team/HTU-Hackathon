"""Grid search over the bot's knobs, scored automatically on grid_eval.json.

    python grid.py            -> full grid (stage 1: context x style x facts, stage 2: temperature x decompose)
    python grid.py --quick    -> stage 1 only
    python grid.py --best     -> only the current best_config.json (quick check after code changes)
    python grid.py --best --file holdout_eval.json   -> best config on the holdout set + spreadsheet

Score per message (0..1):
  answer : fraction of required facts present (numbers, list items, key words); 0 if a known-wrong fact appears;
           minus a repetition penalty (repeated word pairs, e.g. '25 دينار' said twice)
  refuse : 1 if it declines without giving numbers
  clarify: 1 if it asks a question naming >= 2 of the documents
Writes results/grid_<time>.csv, results/grid_<time>_answers.jsonl and best_config.json (main.py loads it).
"""
import copy
import csv
import itertools
import json
import re
import sys
import time
from pathlib import Path

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE.parent))          # main.py is in the project root
import main as M  # noqa: E402
EVAL_FILE = sys.argv[sys.argv.index("--file") + 1] if "--file" in sys.argv else "grid_eval.json"
EVAL = json.loads((HERE / EVAL_FILE).read_text(encoding="utf-8"))

STAGE1 = {"context": ["focused", "full"], "style": ["concise", "detailed"], "facts": [True, False]}
STAGE2 = {"temperature": [0.0, 0.3], "decompose": [True, False]}

REFUSE_RX = re.compile(M.norm(r"عذرا|لا تتوفر|لا املك|لا استطيع|غير متوفر|ليس لدي|sorry|don't have|do not have"))


def repetition(ans):
    w = re.findall(r"\w+", M.norm(ans))
    grams = [tuple(w[i:i + 2]) for i in range(len(w) - 1)]
    return 0.0 if not grams else 1 - len(set(grams)) / len(grams)


def score(msg):
    chk, ans = msg["check"], msg.get("answer", "")
    a = M.norm(ans)
    if chk["kind"] == "refuse":
        return float(bool(REFUSE_RX.search(a)) and not re.search(r"\d", a)), 0.0
    if chk["kind"] == "clarify":
        return float(("؟" in ans or "?" in ans) and len(M.DOC_WORDS.findall(a)) >= 2), 0.0
    if chk.get("polarity"):                       # yes/no questions: the verdict outside quotes must match
        free = re.sub(r"«[^»]*»|\"[^\"]*\"", " ", a)
        m = re.search(r"(?<!\w)(نعم|لا|yes|no)(?!\w)", free)
        if not m or {"yes": "نعم", "no": "لا"}.get(m.group(1), m.group(1)) != M.norm(chk["polarity"]):
            return 0.0, repetition(ans)
    if any(M.norm(x) in a for x in chk["not"]):
        return 0.0, repetition(ans)
    if not chk["must"]:
        return 1.0, repetition(ans)

    def has(alt):
        alt = M.norm(alt)
        return re.search(rf"(?<!\d){alt}(?!\d)", a) if alt.isdigit() else alt in a

    cov = sum(any(has(x) for x in alts) for alts in chk["must"]) / len(chk["must"])
    return cov, repetition(ans)


PROGRESS = {"done": 0, "total": 1, "t0": time.time()}


def _tick(*_):
    """Called after every answered message: prints overall progress + ETA on one line."""
    p = PROGRESS
    p["done"] += 1
    el = time.time() - p["t0"]
    eta = el / p["done"] * (p["total"] - p["done"])
    print(f"\r  {p['done']}/{p['total']} messages  {100 * p['done'] / p['total']:.0f}%  "
          f"elapsed {el / 60:.1f} min  ETA {eta / 60:.1f} min   ", end="", flush=True)


def evaluate(cfg, services, tag):
    M.CONFIG.update(cfg)
    data = copy.deepcopy(EVAL)
    t0 = time.time()
    for _ in M.run_chats(data, services, log=lambda line="", *a: _tick() if str(line).startswith("[") else None):
        pass
    rows, scores, reps = [], [], []
    for chat in data["chats"]:
        for m in chat["messages"]:
            sc, rp = score(m)
            scores.append(sc)
            reps.append(rp)
            rows.append({"setup": tag, "id": m["message_id"], "score": round(sc, 2), "q": m["question"],
                         "expected": chat.get("expected", ""), "evidence": chat.get("evidence", ""),
                         "trap": chat.get("trap", ""), "answer": m.get("answer", ""),
                         "contexts": m.get("contexts", []), "repetition": round(rp, 3)})
    n = len(scores)
    acc = sum(scores) / n
    rep = sum(reps) / n
    overall = acc - 0.5 * rep
    return {"setup": tag, "overall": round(overall, 4), "accuracy": round(acc, 4), "repetition": round(rep, 4),
            "full_marks": sum(s == 1 for s in scores), "n": n, "sec_per_msg": round((time.time() - t0) / n, 2)}, rows


def write_sheet(rows, base):
    """Spreadsheet: question | expected | bot answer | auto score | contexts. xlsx if openpyxl exists, else csv."""
    cols = ["id", "score", "q", "expected", "answer", "contexts", "evidence", "trap"]
    heads = ["ID", "Auto score", "Question", "Expected", "Bot answer", "Contexts returned", "Evidence", "Trap"]
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        wb = Workbook()
        ws = wb.active
        ws.title = "results"
        ws.sheet_view.rightToLeft = False
        ws.append(heads)
        for c in ws[1]:
            c.font = Font(bold=True)
        for r in rows:
            ws.append([r["id"], r["score"], r["q"], r["expected"], r["answer"], "\n".join(r["contexts"]),
                       r["evidence"], r["trap"]])
            fill = "C6EFCE" if r["score"] >= 1 else "FFEB9C" if r["score"] >= 0.5 else "FFC7CE"
            ws.cell(ws.max_row, 2).fill = PatternFill("solid", fgColor=fill)
        for col, w in zip("ABCDEFGH", (11, 10, 45, 45, 60, 60, 28, 28)):
            ws.column_dimensions[col].width = w
        for row in ws.iter_rows(min_row=2):
            for c in row:
                c.alignment = Alignment(wrap_text=True, vertical="top")
        ws.freeze_panes = "C2"
        path = base.with_suffix(".xlsx")
        wb.save(path)
    except ImportError:
        path = base.with_suffix(".csv")
        with open(path, "w", newline="", encoding="utf-8-sig") as fh:         # utf-8-sig: Excel shows Arabic
            w = csv.writer(fh)
            w.writerow(heads)
            for r in rows:
                w.writerow([r[c] if c != "contexts" else " | ".join(r[c]) for c in cols])
    return path


def combos(axes):
    keys = list(axes)
    for vals in itertools.product(*axes.values()):
        yield dict(zip(keys, vals))


def main():
    quick = "--quick" in sys.argv
    only_best = "--best" in sys.argv
    services = M.load_services()
    M.ensure_ollama()
    base = dict(M.CONFIG)
    out = HERE / "results"
    out.mkdir(exist_ok=True)
    stamp = time.strftime("%Y%m%d_%H%M%S")
    results, all_rows = [], []

    def run(cfg):
        tag = " ".join(f"{k}={v}" for k, v in cfg.items())
        r, rows = evaluate(cfg, services, tag)
        results.append((r, dict(cfg)))
        all_rows.extend(rows)
        print(f"\r{tag:<85} overall={r['overall']:.3f} acc={r['accuracy']:.3f} rep={r['repetition']:.3f} "
              f"full={r['full_marks']}/{r['n']}  {r['sec_per_msg']}s/msg", flush=True)

    s1 = list(combos(STAGE1))
    n_msgs = sum(len(c["messages"]) for c in EVAL["chats"])
    if only_best:                                 # confirm the current best_config.json after code changes
        PROGRESS.update(done=0, t0=time.time(), total=n_msgs)
        run(dict(base))
        print("\nStats:", json.dumps(M.STATS))
        name = Path(EVAL_FILE).stem
        f = out / f"best_{name}_{stamp}_answers.jsonl"
        f.write_text("\n".join(json.dumps(r, ensure_ascii=False) for r in all_rows), encoding="utf-8")
        sheet = write_sheet(all_rows, out / f"best_{name}_{stamp}")
        print(f"Answers + scores: {f.relative_to(HERE)}\nSpreadsheet:      {sheet.relative_to(HERE)}")
        return
    PROGRESS.update(done=0, t0=time.time(), total=n_msgs * (len(s1) + (0 if quick else 3)))
    print(f"=== STAGE 1: {len(s1)} setups x {sum(len(c['messages']) for c in EVAL['chats'])} messages ===")
    for c in s1:
        run({**base, **c})
    best = max(results, key=lambda x: x[0]["overall"])[1]
    if not quick:
        s2 = [c for c in combos(STAGE2) if not (c["temperature"] == best["temperature"] and c["decompose"] == best["decompose"])]
        print(f"\n=== STAGE 2: best stage-1 setup x {len(s2)} more (temperature, decompose) ===")
        for c in s2:
            run({**best, **c})
    results.sort(key=lambda x: x[0]["overall"], reverse=True)
    with open(out / f"grid_{stamp}.csv", "w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=list(results[0][0].keys()))
        w.writeheader()
        w.writerows(r for r, _ in results)
    (out / f"grid_{stamp}_answers.jsonl").write_text(
        "\n".join(json.dumps(r, ensure_ascii=False) for r in all_rows), encoding="utf-8")
    print("\nRanking:")
    for r, _ in results:
        print(f"  {r['overall']:.3f}  acc={r['accuracy']:.3f}  rep={r['repetition']:.3f}  {r['setup']}")
    winner = results[0][1]
    (HERE.parent / "best_config.json").write_text(json.dumps(winner, indent=2), encoding="utf-8")
    print(f"\nSaved best_config.json -> main.py uses it automatically:\n  {winner}")
    print(f"Per-question answers + scores: results/grid_{stamp}_answers.jsonl  (search \"score\": 0)")


if __name__ == "__main__":
    main()
